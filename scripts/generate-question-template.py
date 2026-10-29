"""Optional development tool; generated template is committed, no Python needed by users."""
import json
import subprocess
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

headers = ['回合','回合名稱','題目','選項A','選項B','選項C','選項D','正確答案']
questions = json.loads(subprocess.check_output(['node','--input-type=module','-e',"import {QUESTIONS} from './src/questions.js'; console.log(JSON.stringify(QUESTIONS))"]).decode())

def workbook(rows):
    book = Workbook()
    sheet = book.active
    sheet.title = '題庫'
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill('solid',fgColor='11263F')
        cell.font = Font(color='FFFFFF',bold=True,size=12)
    for col,width in zip('ABCDEFGH',[8,24,48,24,24,24,24,14]):
        sheet.column_dimensions[col].width = width
    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height = 45
        for cell in row:
            cell.alignment = Alignment(wrap_text=True,vertical='center')
    for column,values in [('A','1,2,3'),('H','A,B,C,D')]:
        rule = DataValidation(type='list',formula1='"'+values+'"',allow_blank=False)
        rule.showErrorMessage = True
        rule.error = '請使用下拉選單的值'
        sheet.add_data_validation(rule)
        rule.add(column+'2:'+column+'61')
    info = book.create_sheet('填寫說明')
    for text in [
        '將「題庫」工作表的示範題目替換為正式內容。',
        '保留第 1 列八個欄位名稱與「題庫」工作表名稱。',
        '依回合 1、2、3 排列，每回合 1～20 題，同一回合名稱一致。',
        '四個選項皆需填寫且不可重複；正確答案填 A、B、C 或 D。',
        '回合名稱最多 40 字、題目 200 字、各選項 80 字；可填數字選項，不使用公式。',
        '每題答對固定 260 分；檔案不超過 2 MiB。',
        '先預覽確認，再發布；開始遊戲後不能修改當場題庫。',
        '建立新場次會沿用已發布題庫，但玩家與分數重新開始。'
    ]:
        info.append([text])
    info.column_dimensions['A'].width = 95
    return book

rows = [[q['round'],q['roundName'],q['text'],*q['options'],'ABCD'[q['correct']]] for q in questions]
workbook(rows).save('dist/question-bank-template.xlsx')
fixtures = Path('tests/fixtures')
fixtures.mkdir(parents=True,exist_ok=True)
custom = [[r,f'回合{r}',f'正式題目{i+1}','甲','乙','丙','丁','B'] for i,r in enumerate([1,1,2,3])]
workbook(custom).save(fixtures/'bank-valid.xlsx')
invalid = [r.copy() for r in custom]
invalid[1][7] = 'E'
invalid[2][4] = '甲'
workbook(invalid).save(fixtures/'bank-invalid.xlsx')
workbook(custom[:3]).save(fixtures/'bank-missing-round.xlsx')
print('Generated template and local Excel fixtures inside project')

blank = workbook(custom)
blank['題庫']['H2'] = None
blank.save(fixtures/'bank-empty-answer.xlsx')
numeric = workbook(custom)
for column, number in zip('DEFG', [2023, 2024, 2025, 2026]):
    numeric['題庫'][column+'2'] = number
numeric.save(fixtures/'bank-numeric-options.xlsx')
