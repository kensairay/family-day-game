"""Generate a public 24-question rehearsal bank using the existing Excel template.

Run from the repository root with the existing optional openpyxl development tool.
"""
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

# Each tuple contains question, four choices, correct letter, and host explanation.
rounds = [
    ("生活小高手", [
        ("過馬路時，看到哪一種行人號誌才可以開始通行？", ["綠色行人號誌", "紅色行人號誌", "所有號誌熄滅", "只看旁邊的人"], "A", "綠燈時仍要確認來車，注意轉彎車輛。"),
        ("一個星期有幾天？", ["5 天", "7 天", "8 天", "10 天"], "B", "星期一到星期日，共七天。"),
        ("餐前清潔雙手，哪一種做法比較好？", ["只用衣服擦手", "只甩一甩手", "用肥皂和清水洗手", "摸一下桌面"], "C", "用肥皂和清水仔細洗手，再擦乾。"),
        ("吃完香蕉後，哪一個是香蕉皮？", ["紙盒", "鋁罐", "玻璃瓶", "香蕉外面的果皮"], "D", "香蕉皮是香蕉外面的果皮，垃圾分類請依現場規定。"),
        ("在圖書館看書時，哪一種行為最合適？", ["輕聲交談", "大聲唱歌", "追逐奔跑", "開擴音播放影片"], "A", "保持安靜，尊重其他讀者。"),
        ("下雨出門時，哪一樣用品最適合拿來遮雨？", ["太陽眼鏡", "雨傘", "鉛筆", "毛巾架"], "B", "雨傘可以遮雨，使用時注意周圍的人。"),
        ("使用完水龍頭後，怎麼做可以節約用水？", ["讓水一直流", "把水開得更大", "確實關緊水龍頭", "再開另一個水龍頭"], "C", "用完關緊水龍頭，避免浪費。"),
        ("在活動現場與家人走散時，哪一種做法比較安全？", ["獨自離開會場", "跟陌生人去停車場", "躲起來不出聲", "找現場工作人員協助"], "D", "找可辨識的現場工作人員協助聯絡家人。"),
    ]),
    ("動物自然王", [
        ("哪一種動物具有羽毛？", ["麻雀", "兔子", "青蛙", "烏龜"], "A", "鳥類具有羽毛，麻雀是鳥類。"),
        ("蝴蝶長大之前，會經過哪一種幼蟲階段？", ["蝌蚪", "毛毛蟲", "小魚", "蚯蚓"], "B", "蝴蝶的生命週期包含卵、幼蟲、蛹、成蝶。"),
        ("植物通常用哪一個部位從土壤吸收水分？", ["花瓣", "果實", "根", "種皮"], "C", "根能吸收土壤中的水分與礦物質。"),
        ("哪一種動物是哺乳類？", ["金魚", "蝴蝶", "麻雀", "海豚"], "D", "海豚用肺呼吸，並以乳汁哺育幼兒。"),
        ("在一般情況下，水放進冷凍庫結冰後，變成什麼狀態？", ["固態", "氣態", "電漿態", "仍是液態"], "A", "冰是固態的水。"),
        ("下列哪一種動物有八隻腳？", ["螞蟻", "蜘蛛", "蝴蝶", "蚊子"], "B", "蜘蛛有八隻腳，題目中的昆蟲有六隻腳。"),
        ("地球繞著哪一個天體運行？", ["月球", "北極星", "太陽", "火星"], "C", "地球繞太陽公轉，月球繞地球運行。"),
        ("哪一種動物主要用鰓在水中呼吸？", ["貓", "麻雀", "海豚", "金魚"], "D", "金魚用鰓呼吸，海豚則需要浮出水面用肺呼吸。"),
    ]),
    ("歡樂動腦題", [
        ("桌上有 3 顆蘋果，再放上 2 顆，一共有幾顆？", ["5 顆", "4 顆", "6 顆", "7 顆"], "A", "3 加 2 等於 5。"),
        ("依照規律：2、4、6，下一個數字是多少？", ["7", "8", "9", "10"], "B", "每次加 2，下一個是 8。"),
        ("一個三角形有幾條邊？", ["1 條", "2 條", "3 條", "4 條"], "C", "三角形有三條邊。"),
        ("現在是下午 2 點，再過 1 小時是幾點？", ["下午 1 點", "下午 4 點", "下午 5 點", "下午 3 點"], "D", "下午 2 點加一小時，是下午 3 點。"),
        ("10 顆糖果平均分給 5 個人，每人可以拿到幾顆？", ["2 顆", "3 顆", "4 顆", "5 顆"], "A", "10 除以 5 等於 2。"),
        ("面向前方時，向右轉兩次，每次轉 90 度，最後面向哪裡？", ["原本的右方", "原本的後方", "原本的左方", "原本的前方"], "B", "兩次 90 度合計 180 度，面向原本的後方。"),
        ("跑步比賽中，小明超過了原本的第二名，此時小明是第幾名？", ["第一名", "第三名", "第二名", "第四名"], "C", "超過第二名後，接替第二名的位置。"),
        ("每人都和其他人各握手一次，3 個人總共握手幾次？", ["1 次", "2 次", "6 次", "3 次"], "D", "三組配對分別是甲乙、甲丙、乙丙，不重複計算。"),
    ]),
]

book = load_workbook("dist/question-bank-template.xlsx")
sheet = book["題庫"]
answer_sheet = book.create_sheet("主持人答案")
answer_sheet.append(["總題號", "回合", "回合題號", "題目", "正確答案", "答案內容", "解說"])
row_number = 2
for round_number, (name, questions) in enumerate(rounds, start=1):
    for number, (question, options, letter, explanation) in enumerate(questions, start=1):
        for column, value in enumerate([round_number, name, question, *options, letter], start=1):
            sheet.cell(row_number, column, value)
        sheet.row_dimensions[row_number].height = 60
        for cell in sheet[row_number]:
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.fill = PatternFill("solid", fgColor="EEF5FF" if round_number % 2 else "F0FAF2")
        answer_sheet.append([row_number - 1, round_number, number, question, letter, options["ABCD".index(letter)], explanation])
        row_number += 1
sheet.auto_filter.ref = sheet.dimensions
answer_sheet.freeze_panes = "A2"
answer_sheet.auto_filter.ref = answer_sheet.dimensions
for cell in answer_sheet[1]:
    cell.fill = PatternFill("solid", fgColor="11263F")
    cell.font = Font(color="FFFFFF", bold=True)
for column, width in zip("ABCDEFG", [10, 8, 12, 60, 12, 28, 65]):
    answer_sheet.column_dimensions[column].width = width
for row in answer_sheet.iter_rows(min_row=2):
    answer_sheet.row_dimensions[row[0].row].height = 48
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="center")
info = book["填寫說明"]
for text in [
    "本檔為公開測試題庫，共三回合、每回合八題；適合家庭日流程彩排，並非正式競賽難度設計。",
    "題庫工作表可直接匯入；主持人答案工作表供核對與解說，不會匯入遊戲。",
    "每回合滿分 2080 分，三回合全對共 6240 分；不計答題速度加成。",
    "請勿在投影幕或玩家畫面公開本檔：題庫與主持人答案工作表均包含正解。",
]:
    info.append([text])
output = Path("dist/question-bank-test-24.xlsx")
book.save(output)
print(f"Generated {output}: 3 rounds, 8 questions each, 24 questions total")
